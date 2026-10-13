#!/usr/bin/env python3
"""Generate a review sheet (xlsx) for one form, for the MSF team.

The sheet lists every question as the system configures it today (answers,
mandatory, when it is shown, limits), in plain French when a translation
exists. MSF only marks each line as correct or not and adds scenarios.
The returned file is the input for the lime-test-case skill.

Usage:
  python3 scripts/form-test-sheet.py <form.json> [--lang fr] [--out file.xlsx]
"""
import argparse
import datetime
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF4C2")
HEADER = PatternFill("solid", fgColor="1F4E79")
SECTION = PatternFill("solid", fgColor="DDEBF7")
GREY = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

RENDERING = {
    "radio": "Choix unique",
    "select": "Choix unique",
    "content-switcher": "Choix unique",
    "multiCheckbox": "Choix multiple",
    "checkbox": "Choix multiple",
    "number": "Nombre",
    "text": "Texte libre",
    "textarea": "Texte libre",
    "date": "Date",
    "datetime": "Date et heure",
    "toggle": "Oui / Non",
    "drug": "Médicament",
    "problem": "Diagnostic",
    "file": "Fichier / photo",
    "workspace-launcher": "Bouton",
    "ui-select-extended": "Liste de recherche",
}
SKIP_RENDERINGS = {"markdown", "fixed-value"}
PATIENT_FIELDS = {"age": "l'âge du patient", "sex": "le sexe du patient"}
LONG_LIST = 15
UNKNOWN_ANSWER = "« réponse introuvable dans la question, à vérifier »"


class Form:
    def __init__(self, path, lang):
        with open(path, encoding="utf-8") as f:
            self.data = json.load(f)
        self.tr = self._translations(path, lang)
        self.by_id = {}
        self.answer_labels = {}
        for _, q in self.walk():
            if q.get("id"):
                self.by_id[q["id"]] = q
            for a in q.get("questionOptions", {}).get("answers", []):
                self.answer_labels.setdefault(a.get("concept"), a.get("label", ""))

    @staticmethod
    def _translations(path, lang):
        base = os.path.splitext(os.path.basename(path))[0]
        tr_dir = os.path.join(os.path.dirname(os.path.dirname(path)), "ampathformstranslations")
        tr_file = os.path.join(tr_dir, f"{base}_translations_{lang}.json")
        if os.path.exists(tr_file):
            with open(tr_file, encoding="utf-8") as f:
                return json.load(f).get("translations", {})
        return {}

    def t(self, text):
        return self.tr.get(text, text) if text else ""

    def walk(self):
        """Yields (section label, question), flattening groups."""
        def rec(section, questions):
            for q in questions:
                yield section, q
                yield from rec(section, q.get("questions", []))
        for page in self.data.get("pages", []):
            for section in page.get("sections", []):
                yield from rec(section.get("label", ""), section.get("questions", []))

    def qlabel(self, qid):
        if qid in PATIENT_FIELDS:
            return PATIENT_FIELDS[qid]
        q = self.by_id.get(qid)
        return f"« {self.t(q.get('label', qid))} »" if q else qid

    def alabel(self, qid, concept):
        q = self.by_id.get(qid, {})
        for a in q.get("questionOptions", {}).get("answers", []):
            if a.get("concept") == concept:
                return f"« {self.t(a.get('label'))} »"
        label = self.answer_labels.get(concept)
        return f"« {self.t(label)} »" if label else UNKNOWN_ANSWER


# --- skip logic in plain French -------------------------------------------

# Each pattern gives (what the clause says, its opposite) in plain French.
CLAUSE_PATTERNS = [
    (r"^!\s*includes\(\s*(\w+)\s*,\s*'([^']+)'\s*\)$", "{q} ne contient pas {a}", "{q} contient {a}"),
    (r"^includes\(\s*(\w+)\s*,\s*'([^']+)'\s*\)$", "{q} contient {a}", "{q} ne contient pas {a}"),
    (r"^(\w+)\s*!==?\s*'([^']+)'$", "{q} ≠ {a}", "{q} = {a}"),
    (r"^(\w+)\s*===?\s*'([^']+)'$", "{q} = {a}", "{q} ≠ {a}"),
    (r"^!\s*isEmpty\(\s*(\w+)\s*\)$", "{q} est renseigné", "{q} est vide"),
    (r"^isEmpty\(\s*(\w+)\s*\)$", "{q} est vide", "{q} est renseigné"),
]
NUMERIC = re.compile(r"^(\w+)\s*(<=|>=|<|>|!==?|===?)\s*(-?\d+(?:\.\d+)?)$")
SAYS = {"<": "<", "<=": "≤", ">": ">", ">=": "≥", "!==": "≠", "!=": "≠", "===": "=", "==": "="}
OPPOSITE = {"<": "≥", "<=": ">", ">": "≤", ">=": "<", "!==": "=", "!=": "=", "===": "≠", "==": "≠"}


def strip_parens(s):
    s = s.strip()
    while s.startswith("(") and s.endswith(")") and balanced(s[1:-1]):
        s = s[1:-1].strip()
    return s


def balanced(s):
    depth = 0
    for c in s:
        depth += c == "("
        depth -= c == ")"
        if depth < 0:
            return False
    return depth == 0


def split_top(s, op):
    parts, depth, cur, i = [], 0, "", 0
    while i < len(s):
        c = s[i]
        depth += c == "("
        depth -= c == ")"
        if depth == 0 and s.startswith(op, i):
            parts.append(cur)
            cur, i = "", i + len(op)
            continue
        cur += c
        i += 1
    parts.append(cur)
    return [strip_parens(p) for p in parts]


def clause(form, text):
    """Returns (what the clause says, its opposite), or None if not understood."""
    m = NUMERIC.match(text)
    if m:
        q = form.qlabel(m.group(1))
        return f"{q} {SAYS[m.group(2)]} {m.group(3)}", f"{q} {OPPOSITE[m.group(2)]} {m.group(3)}"
    for regex, says, opposite in CLAUSE_PATTERNS:
        m = re.match(regex, text)
        if m:
            qid = m.group(1)
            a = form.alabel(qid, m.group(2)) if m.lastindex and m.lastindex > 1 else ""
            return says.format(q=form.qlabel(qid), a=a), opposite.format(q=form.qlabel(qid), a=a)
    return None


def visibility(form, expr):
    """hideWhenExpression is "A || B || (C && D)"; the field shows when
    not A AND not B AND (not C OR not D)."""
    if not expr:
        return "Toujours"
    shown = []
    for part in split_top(strip_parens(expr), "||"):
        clauses = [clause(form, c) for c in split_top(part, "&&")]
        if any(c is None for c in clauses):
            return "Sous condition (à vérifier avec l'équipe technique)"
        opposites = [opp for _, opp in clauses]
        shown.append(opposites[0] if len(opposites) == 1 else "(" + " OU ".join(opposites) + ")")
    return "Seulement si " + " ET ".join(shown)


# --- rows ----------------------------------------------------------------

def is_required(q):
    r = q.get("required")
    if isinstance(r, str):
        return r.lower() == "true"
    if isinstance(r, dict):
        return True
    return bool(r)


def rules(form, q):
    qo = q.get("questionOptions", {})
    out = []
    lo, hi = qo.get("min"), qo.get("max")
    if lo not in (None, "") and hi not in (None, ""):
        out.append(f"Entre {lo} et {hi}")
    elif lo not in (None, ""):
        out.append(f"Minimum {lo}")
    elif hi not in (None, ""):
        out.append(f"Maximum {hi}")
    for v in q.get("validators", []) or []:
        if v.get("type") == "date" and not v.get("allowFutureDates"):
            out.append("Pas de date future")
        elif v.get("message"):
            out.append(form.t(v["message"]))
    return " ; ".join(out)


def landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def build(form_path, lang, out):
    form = Form(form_path, lang)
    meta = form.data
    title = meta.get("name", os.path.basename(form_path))
    wb = Workbook()

    # Lisez-moi
    ws = wb.active
    ws.title = "Lisez-moi"
    lines = [
        (f"Fiche de vérification : {title}", Font(name=FONT, size=16, bold=True, color="1F4E79")),
        (meta.get("description", ""), Font(name=FONT, italic=True, color="595959")),
        ("", None),
        ("Merci de nous aider à vérifier ce formulaire. Cela prend environ 20 minutes.", Font(name=FONT, size=12)),
        ("", None),
        ("1.  Onglet « Questions » : chaque ligne montre ce que fait le système aujourd'hui.", Font(name=FONT, size=11)),
        ("      Dans la colonne jaune « Correct ? », choisissez Oui ou Non.", Font(name=FONT, size=11)),
        ("      Si Non, écrivez simplement ce qu'il faudrait dans la colonne suivante.", Font(name=FONT, size=11)),
        ("2.  Onglet « Scénarios » : ajoutez les situations importantes à tester,", Font(name=FONT, size=11)),
        ("      sous la forme « Si je fais... je dois voir... ». La première ligne est un exemple.", Font(name=FONT, size=11)),
        ("3.  Renvoyez-nous le fichier. Nous préparons les tests à partir de vos réponses.", Font(name=FONT, size=11)),
        ("", None),
        ("Légende", Font(name=FONT, bold=True)),
        ("Cellules à remplir", Font(name=FONT)),
        ("Ne pas modifier", Font(name=FONT)),
        ("", None),
        ("N'utilisez jamais de données réelles de patients dans les exemples.", Font(name=FONT, bold=True, color="C00000")),
        ("", None),
        (f"Formulaire : {os.path.basename(form_path)}, version {meta.get('version', '?')}. "
         f"Fiche générée le {datetime.date.today().isoformat()}"
         + ("" if form.tr else " (pas de traduction française : libellés en anglais)") + ".",
         Font(name=FONT, size=9, color="808080")),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text)
        if font:
            ws.cell(row=i, column=1).font = font
    ws.cell(row=14, column=1).fill = YELLOW
    ws.cell(row=15, column=1).fill = GREY
    ws.column_dimensions["A"].width = 95
    ws.sheet_view.showGridLines = False

    # Questions
    ws = wb.create_sheet("Questions")
    headers = [
        ("N°", 5), ("Question", 42), ("Type de réponse", 15), ("Réponses possibles", 40),
        ("Obligatoire ?", 12), ("Quand est-elle affichée ?", 40), ("Limites / règles", 26),
        ("Correct ?", 12), ("Si non, que faut-il ?", 40), ("Remarque ou exemple (facultatif)", 32),
        ("ID technique", 30),
    ]
    for c, (h, w) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    ws.cell(row=1, column=8).comment = Comment("Choisissez Oui, Non ou Je ne sais pas.", "LIME")

    dv = DataValidation(type="list", formula1='"Oui,Non,Je ne sais pas"', allow_blank=True)
    ws.add_data_validation(dv)

    row, n, last_section = 2, 0, None
    long_lists = []
    for section, q in form.walk():
        qo = q.get("questionOptions", {})
        rendering = qo.get("rendering", "")
        if rendering in SKIP_RENDERINGS or not q.get("label"):
            continue
        if section != last_section:
            ws.cell(row=row, column=1, value=form.t(section))
            for c in range(1, len(headers) + 1):
                ws.cell(row=row, column=c).fill = SECTION
                ws.cell(row=row, column=c).border = BORDER
            ws.cell(row=row, column=1).font = Font(name=FONT, bold=True, color="1F4E79")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            row += 1
            last_section = section
        n += 1
        is_group = rendering == "group" or q.get("type") == "obsGroup"
        calc = "calculate" in qo
        kind = "Groupe de questions" if is_group else ("Calculé automatiquement" if calc else RENDERING.get(rendering, rendering))
        labels = [form.t(a.get("label", "")) for a in qo.get("answers", [])]
        if len(labels) > LONG_LIST:
            long_lists.append((n, form.t(q.get("label")), labels))
            answers = " ; ".join(labels[:10]) + f" ... et {len(labels) - 10} autres (voir l'onglet « Listes longues »)"
        else:
            answers = " ; ".join(labels)
        values = [
            n,
            form.t(q.get("label")),
            kind,
            answers,
            "" if is_group or calc else ("Oui" if is_required(q) else "Non"),
            visibility(form, q.get("hide", {}).get("hideWhenExpression", "")),
            rules(form, q),
            None,
            None,
            None,
            q.get("id", ""),
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = Font(name=FONT, size=10, bold=is_group, color="808080" if c == 11 else "000000")
            cell.alignment = WRAP
            cell.border = BORDER
            cell.fill = YELLOW if c in (8, 9, 10) else GREY
        dv.add(f"H{row}")
        row += 1

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(headers)).column_letter}{row - 1}"
    ws.column_dimensions["K"].hidden = True
    landscape(ws)

    # Scénarios
    ws = wb.create_sheet("Scénarios")
    sheaders = [("N°", 5), ("Pour qui ? (rôle, site)", 26), ("Si je fais...", 50), ("Je dois voir...", 50), ("Important ?", 14)]
    for c, (h, w) in enumerate(sheaders, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    example = ["Ex.", "Clinicien, tous les sites",
               "Exemple : je réponds « Oui » à la question X",
               "Exemple : la question Y apparaît et doit être remplie", "Critique"]
    for c, v in enumerate(example, start=1):
        cell = ws.cell(row=2, column=c, value=v)
        cell.font = Font(name=FONT, size=10, italic=True, color="808080")
        cell.alignment = WRAP
        cell.border = BORDER
    idv = DataValidation(type="list", formula1='"Critique,Normal"', allow_blank=True)
    ws.add_data_validation(idv)
    for r in range(3, 23):
        ws.cell(row=r, column=1, value=r - 2).font = Font(name=FONT, size=10)
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = WRAP
            if c > 1:
                cell.fill = YELLOW
                cell.font = Font(name=FONT, size=10)
        idv.add(f"E{r}")
    ws.freeze_panes = "A2"
    landscape(ws)

    if long_lists:
        ws = wb.create_sheet("Listes longues")
        for c, (h, w) in enumerate([("N° question", 11), ("Question", 40), ("Réponse possible", 60), ("Remarque (facultatif)", 40)], start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = Font(name=FONT, bold=True, color="FFFFFF")
            cell.fill = HEADER
            ws.column_dimensions[cell.column_letter].width = w
        r = 2
        for qn, qlabel, labels in long_lists:
            for label in labels:
                for c, v in enumerate([qn, qlabel, label, None], start=1):
                    cell = ws.cell(row=r, column=c, value=v)
                    cell.font = Font(name=FONT, size=10)
                    cell.border = BORDER
                    cell.fill = YELLOW if c == 4 else GREY
                r += 1
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:D{r - 1}"

    wb.active = 0
    wb.save(out)
    return n


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("form")
    p.add_argument("--lang", default="fr")
    p.add_argument("--out")
    args = p.parse_args()
    out = args.out or os.path.splitext(os.path.basename(args.form))[0] + "_verification.xlsx"
    n = build(args.form, args.lang, out)
    print(f"{out}: {n} questions")


if __name__ == "__main__":
    sys.exit(main())
